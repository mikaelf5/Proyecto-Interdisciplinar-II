"""Genera el archivo del reporte: Excel con openpyxl, PDF con reportlab."""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .servicios import generar_datos

EXCEL = "excel"
PDF = "pdf"
FORMATOS = (EXCEL, PDF)

CONTENT_TYPES = {
    EXCEL: "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    PDF: "application/pdf",
}
EXTENSIONES = {EXCEL: "xlsx", PDF: "pdf"}


def a_excel(datos):
    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"
    ws["A1"] = datos.titulo
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = datos.subtitulo
    fila = 4
    for etiqueta, valor in datos.resumen:
        ws.cell(row=fila, column=1, value=etiqueta).font = Font(bold=True)
        ws.cell(row=fila, column=2, value=valor)
        fila += 1
    fila += 1
    cabecera = PatternFill("solid", fgColor="343A40")
    for col, nombre in enumerate(datos.columnas, start=1):
        celda = ws.cell(row=fila, column=col, value=nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = cabecera
    for datos_fila in datos.filas:
        fila += 1
        for col, valor in enumerate(datos_fila, start=1):
            ws.cell(row=fila, column=col, value=valor).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
    for col in range(1, max(len(datos.columnas), 2) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 38 if col in (2, 3, 4) else 22
    salida = BytesIO()
    wb.save(salida)
    return salida.getvalue()


def a_pdf(datos):
    estilos = getSampleStyleSheet()
    celda = ParagraphStyle("celda", parent=estilos["BodyText"], fontSize=7.5, leading=9)
    cabecera = ParagraphStyle("cab", parent=celda, textColor=colors.white, fontName="Helvetica-Bold")

    salida = BytesIO()
    doc = SimpleDocTemplate(
        salida, pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm, topMargin=1.5 * cm, bottomMargin=1.5 * cm,
        title=datos.titulo,
    )
    elementos = [
        Paragraph(datos.titulo, estilos["Title"]),
        Paragraph(datos.subtitulo, estilos["Normal"]),
        Spacer(1, 0.5 * cm),
    ]
    resumen = Table(
        [[Paragraph(f"<b>{e}</b>", celda), Paragraph(v, celda)] for e, v in datos.resumen],
        colWidths=[6 * cm, 4 * cm], hAlign="LEFT",
    )
    resumen.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey)]))
    elementos += [resumen, Spacer(1, 0.6 * cm)]

    if datos.filas:
        filas = [[Paragraph(c, cabecera) for c in datos.columnas]]
        filas += [[Paragraph(str(v), celda) for v in f] for f in datos.filas]
        ancho = (landscape(A4)[0] - 3 * cm) / len(datos.columnas)
        tabla = Table(filas, colWidths=[ancho] * len(datos.columnas), repeatRows=1)
        tabla.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#343a40")),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.lightgrey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f4f4f6")]),
        ]))
        elementos.append(tabla)
    else:
        elementos.append(Paragraph("No hay registros en este periodo.", estilos["Normal"]))
    doc.build(elementos)
    return salida.getvalue()


def exportar(reporte, formato):
    """Devuelve (bytes, content_type, nombre_de_archivo). Lanza ReporteNoDisponible."""
    if formato not in FORMATOS:
        raise ValueError(f"Formato no soportado: {formato}")
    datos = generar_datos(reporte)
    contenido = a_excel(datos) if formato == EXCEL else a_pdf(datos)
    nombre = (
        f"reporte_{reporte.tipo.lower()}_{reporte.periodo_inicio:%Y%m%d}"
        f"_{reporte.periodo_fin:%Y%m%d}.{EXTENSIONES[formato]}"
    )
    return contenido, CONTENT_TYPES[formato], nombre
